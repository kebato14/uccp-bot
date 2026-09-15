"""Организационные заявки УЦЦП: отдельное хранение и раздельная статистика.

Запуск: .venv/bin/python tests/org_check.py
"""
from __future__ import annotations

import asyncio
import datetime as dt
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ["BOT_TOKEN"] = "test:token"
os.environ["DB_URL"] = "sqlite+aiosqlite:///" + os.path.join(tempfile.mkdtemp(), "org.db")

from sqlalchemy import select  # noqa: E402

from uccp_bot.db.base import SessionMaker, init_db  # noqa: E402
from uccp_bot.db.models import (  # noqa: E402
    Brand, Category, OrgAttachment, OrgRequest, OrgStatus, OrgType, Outlet, Priority,
    Request, RoleCode, Status, User, UserStatus,
)
from uccp_bot.db.seed import seed  # noqa: E402
from uccp_bot.services import excel, gsheets, org, reports  # noqa: E402
from uccp_bot.services.numbering import (  # noqa: E402
    new_external_id, next_org_number, next_request_number,
)
from uccp_bot.utils import now_local, utcnow  # noqa: E402


def check(condition: bool, label: str) -> None:
    print(("✅ " if condition else "❌ ") + label)
    if not condition:
        raise AssertionError(label)


async def main() -> None:
    await init_db()
    async with SessionMaker() as session:
        await seed(session)

        moose = await session.scalar(select(Brand).where(Brand.name == "Moose Café"))
        hd = await session.scalar(select(Brand).where(Brand.name == "Hotdogger"))
        opera = await session.scalar(select(Outlet).where(Outlet.name == "Moose Opera"))
        tcell = await session.scalar(select(Outlet).where(Outlet.name == "Moose TCell"))
        electric = await session.scalar(select(Category).where(Category.code == "electric"))
        nazrullo = await session.scalar(select(User).where(User.full_name == "Назрулло"))

        def make(tg, name, role, brand=None, outlet=None):
            u = User(tg_id=tg, full_name=name, last_name=name, role_code=role,
                     brand_id=brand.id if brand else None,
                     outlet_id=outlet.id if outlet else None,
                     status=UserStatus.ACTIVE, is_registered=True)
            u.sync_flags()
            session.add(u)
            return u

        admin = make(9001, "Админ системы", RoleCode.ADMIN)
        ops = make(9002, "Опердир Moose", RoleCode.OPS_DIRECTOR, moose)
        outlet_admin = make(9003, "Админ Оперы", RoleCode.OUTLET_ADMIN, moose, opera)
        staff = make(9004, "Сотрудник Оперы", RoleCode.STAFF, moose, opera)
        uccp = make(9005, "Сотрудник УЦЦП", RoleCode.UCCP_STAFF)
        await session.flush()

        print("=== Кто может обращаться в УЦЦП ===")
        check(org.can_create(outlet_admin), "Администратор точки может обратиться в УЦЦП")
        check(org.can_create(ops) and org.can_create(admin),
              "Руководители и администратор системы тоже")
        check(not org.can_create(staff), "Рядовой сотрудник обращение не создаёт")
        check(not org.can_create(nazrullo), "Мастер обращение не создаёт")
        check(not org.can_create(uccp),
              "Сотрудник УЦЦП обращения обрабатывает, а не подаёт от точки")

        print("\n=== Кто обрабатывает ===")
        check(org.can_process(uccp) and org.can_process(admin),
              "Обрабатывают только сотрудники УЦЦП и администратор")
        check(not org.can_process(outlet_admin), "Управляющий точки не обрабатывает обращения")
        check(not org.can_process(ops), "Операционный директор не обрабатывает обращения")
        check(not org.can_process(nazrullo), "Мастер к обращениям доступа не имеет")
        staff_list = await org.uccp_staff(session)
        check({u.id for u in staff_list} == {uccp.id, admin.id},
              "Новые обращения уходят только в УЦЦП")
        check(nazrullo.id not in {u.id for u in staff_list},
              "Сантехники и электрики в этот поток не попадают")

        # --- ремонтная заявка ---
        tech = Request(
            number=await next_request_number(session), external_id=new_external_id(),
            author_id=staff.id, brand_id=moose.id, outlet_id=opera.id,
            category_id=electric.id, executor_id=nazrullo.id,
            description="Не работает розетка", priority=Priority.HIGH,
            status=Status.CLOSED, due_at=utcnow(), work_cost=250, material_cost=120.5,
            closed_at=utcnow(),
        )
        session.add(tech)
        await session.flush()

        # --- организационные заявки ---
        async def add_org(author, assignee=None, outlet=None, object_text=None, brand=None,
                          status=OrgStatus.NEW, cost=None, overdue=False,
                          rtype=OrgType.TASK, quantity=None):
            req = OrgRequest(
                number=await next_org_number(session), external_id=new_external_id(),
                outlet_id=outlet.id if outlet else None,
                brand_id=(outlet.brand_id if outlet else (brand.id if brand else None)),
                object_text=object_text,
                request_type=rtype,
                description="Нужны стаканы 400 мл" if rtype == OrgType.INVENTORY
                else "Подготовить документы по объекту",
                quantity=quantity,
                priority=Priority.MEDIUM, due_at=utcnow() + dt.timedelta(days=2),
                author_id=author.id, author_position=author.position_text,
                assignee_id=assignee.id if assignee else None, status=status,
                cost=cost, is_overdue=overdue,
                closed_at=utcnow() if status == OrgStatus.CLOSED else None,
            )
            session.add(req)
            await session.flush()
            await org.log(session, req, "created", user=author, new_status=OrgStatus.NEW)
            return req

        o1 = await add_org(outlet_admin, uccp, outlet=opera, status=OrgStatus.CLOSED, cost=300,
                           rtype=OrgType.INVENTORY, quantity="200 шт")
        o2 = await add_org(ops, uccp, outlet=tcell, overdue=True)
        o3 = await add_org(admin, uccp, object_text="Офис УЦЦП")
        session.add(OrgAttachment(request_id=o1.id, file_id="F1", media_type="photo",
                                  stage="request"))
        await session.commit()
        o1 = await org.load(session, o1.id)

        print("\n=== Раздельное хранение ===")
        check(o1.number.startswith("ORG-"), f"Своя нумерация: {o1.number}")
        check(tech.number.startswith("REQ-"), f"У ремонтных своя: {tech.number}")
        check(OrgRequest.__tablename__ == "org_requests", "Отдельная таблица в БД")
        tech_count = len((await session.scalars(select(Request))).all())
        org_count = len((await session.scalars(select(OrgRequest))).all())
        check(tech_count == 1 and org_count == 3,
              f"Данные не смешаны: ремонтных {tech_count}, организационных {org_count}")
        check(o3.object_name == "Офис УЦЦП", "Объект вне торговых точек поддерживается")
        check(o1.author_id == outlet_admin.id, "Автор подставлен из профиля")
        check(o1.request_type == OrgType.INVENTORY and o1.quantity == "200 шт",
              "Тип обращения и количество сохранены")

        print("\n=== Карточка по образцу ===")
        card = org.render_card(o1)
        for field in ("ЗАЯВКА В УЦЦП", "Бренд:", "Точка:", "Автор:", "Должность:",
                      "Тип:", "Запрос:", "Количество:", "Срок:", "Фото/файл:"):
            check(field in card, f"В карточке есть «{field}»")
        check("приложено" in card, "Видно, что файл приложен")

        print("\n=== Видимость ===")
        check(org.can_view(uccp, o1) and org.can_view(uccp, o3),
              "Сотрудник УЦЦП видит все обращения")
        check(org.can_view(outlet_admin, o1), "Автор видит своё обращение")
        check(not org.can_view(outlet_admin, o3), "Чужое обращение автору не видно")
        check(not org.can_view(staff, o1), "Посторонний обращение не видит")
        check(not org.can_view(nazrullo, o1), "Мастеру обращения точек не видны")
        check(org.can_close(outlet_admin, o1), "Закрывает автор обращения")
        check(org.can_close(uccp, o1), "Либо сотрудник УЦЦП")
        check(not org.can_close(staff, o2), "Посторонний закрыть не может")

        print("\n=== Цепочка обработки ===")
        chain = [OrgStatus.NEW, OrgStatus.ACCEPTED, OrgStatus.ASSIGNED,
                 OrgStatus.IN_PROGRESS, OrgStatus.DONE, OrgStatus.CLOSED]
        titles = [OrgStatus.title(c) for c in chain]
        check("Принята УЦЦП" in titles[1] and "Назначен ответственный" in titles[2],
              "Статусы: Новая → Принята УЦЦП → Назначен ответственный → …")
        check(OrgStatus.CLARIFY in OrgStatus.OPEN,
              "Возврат на уточнение — рабочий статус, обращение не теряется")

        print("\n=== Раздельная статистика ===")
        today = now_local().date()
        f = reports.ReportFilters(date_from=today, date_to=today, period_title="за сегодня")
        tech_report = await reports.build_report(session, f)
        org_report = await reports.build_org_report(session, f)

        check(tech_report.total == 1, f"В отчёте ремонтных только ремонтные: {tech_report.total}")
        check(abs(tech_report.total_cost - 370.5) < 0.01,
              f"Расходы ремонтных: {tech_report.total_cost}")
        check(org_report.total == 3, f"В отчёте УЦЦП только обращения точек: {org_report.total}")
        check(abs(org_report.total_cost - 300) < 0.01, f"Затраты УЦЦП: {org_report.total_cost}")
        check(org_report.closed == 1 and org_report.overdue == 1,
              "Свои статусы и просрочка считаются отдельно")
        check("Moose Opera" in org_report.by_object and "Офис УЦЦП" in org_report.by_object,
              "Разрез по объектам")
        check("Сотрудник УЦЦП" in org_report.by_assignee, "Разрез по ответственным УЦЦП")
        check(any("инвентарь" in k.lower() for k in org_report.by_type),
              "Разрез по типам обращений")

        org_text = reports.render_org_report(org_report, f)
        check("Заявки в УЦЦП" in org_text, "Отчёт УЦЦП формируется")
        check("REQ-" not in org_text, "В отчёте УЦЦП нет ремонтных заявок")

        print("\n=== Сводный отчёт администратора ===")
        combined = reports.render_combined_report(tech_report, org_report, f)
        check("Технические и ремонтные заявки" in combined and
              "Заявки в УЦЦП" in combined,
              "Оба модуля показаны раздельными блоками")
        check("Заявок: <b>4</b>" in combined, "Сводный итог по количеству")
        check("670" in combined.replace(",", "."), "Сводный итог по расходам: 370,5 + 300")

        print("\n=== Выгрузки ===")
        org_rows = await reports.fetch_org_requests(session, f)
        tech_rows = await reports.fetch_requests(session, f)
        path_org = excel.export_org_requests(org_rows)
        path_all = excel.export_combined(tech_rows, org_rows)
        check(os.path.exists(path_org), "Excel по заявкам УЦЦП выгружен")
        from openpyxl import load_workbook
        names = load_workbook(path_all).sheetnames
        check(names == ["Ремонтные", "Организационные"],
              f"Сводный Excel: два отдельных листа {names}")

        print("\n=== Google Sheets ===")
        payload = gsheets.build_org_month_payload(org_rows, today.year, today.month)
        flat = "\n".join(" | ".join(str(c) for c in row) for row in payload)
        check(payload[0][0].startswith("Заявки в УЦЦП за"), "Заголовок листа УЦЦП")
        check(payload[4] == gsheets.ORG_DETAIL_HEADERS,
              "Шапка: тип, запрос, количество, автор, ответственный, срок")
        check("ЗАТРАТЫ ПО ОБЪЕКТАМ" in flat and "ЗАТРАТЫ ПО ОТВЕТСТВЕННЫМ" in flat,
              "Сводная часть листа УЦЦП")
        check(gsheets.org_worksheet_title(2026, 9) == "2026-09-УЦЦП",
              "Отдельный лист месяца для УЦЦП")
        check("REQ-" not in flat, "На листе УЦЦП нет ремонтных заявок")

    print("\n🎉 Канал «Заявка в УЦЦП» работает и отделён от заявок мастерам.")


if __name__ == "__main__":
    asyncio.run(main())
