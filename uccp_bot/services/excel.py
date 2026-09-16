"""Выгрузка заявок в Excel (п.36).

Функции export_* делают синхронную работу (openpyxl), поэтому вызывать их
из обработчиков нужно через `run_export` — иначе на время формирования файла
бот перестаёт отвечать всем остальным пользователям.
"""
from __future__ import annotations

import os
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import config
from ..db.models import Request, Status
from ..utils import fmt_dt, local_from_utc, now_local

HEADERS = [
    ("№ заявки", 18),
    ("Дата создания", 18),
    ("Бренд", 16),
    ("Точка", 22),
    ("Инициатор", 22),
    ("Категория", 22),
    ("Описание", 50),
    ("Исполнитель", 22),
    ("Приоритет", 14),
    ("Срок", 18),
    ("Статус", 24),
    ("Дата выполнения", 18),
    ("Дата закрытия", 18),
    ("Стоимость работ", 16),
    ("Стоимость материалов", 18),
    ("Итого", 14),
    ("Комментарий исполнителя", 40),
    ("Комментарий менеджера", 40),
    ("Просрочена", 12),
]


def _write_tech_sheet(ws, requests) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")
    for col, (name, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    for row_idx, req in enumerate(requests, start=2):
        values = [
            req.number,
            fmt_dt(req.created_at),
            req.brand.name if req.brand else "",
            req.outlet.name if req.outlet else "",
            req.author.full_name if req.author else "",
            req.category.name if req.category else "",
            req.description,
            req.executor.full_name if req.executor else "не назначен",
            req.priority_title,
            fmt_dt(req.due_at),
            Status.title(req.status),
            fmt_dt(req.done_at),
            fmt_dt(req.closed_at),
            "фикс. оплата" if req.cost_exempt else float(req.work_cost or 0),
            "фикс. оплата" if req.cost_exempt else float(req.material_cost or 0),
            "фикс. оплата" if req.cost_exempt else float(req.total_cost),
            req.executor_comment or "",
            req.manager_comment or "",
            "да" if req.is_overdue else "нет",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            if col in (7, 17, 18):
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    last = len(requests) + 2
    ws.cell(row=last, column=13, value="ИТОГО:").font = Font(bold=True)
    for col in (14, 15, 16):
        letter = get_column_letter(col)
        cell = ws.cell(row=last, column=col, value=f"=SUM({letter}2:{letter}{last - 1})")
        cell.font = Font(bold=True)


def export_requests(requests: List[Request], title: str = "Заявки") -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    _write_tech_sheet(ws, requests)

    os.makedirs(config.export_dir, exist_ok=True)
    stamp = now_local().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(config.export_dir, f"uccp_{stamp}.xlsx")
    wb.save(path)
    return path


ORG_HEADERS = [
    ("№ заявки", 18),
    ("Дата создания", 18),
    ("Точка / объект", 26),
    ("Бренд", 16),
    ("Тип обращения", 26),
    ("Запрос", 50),
    ("Количество", 16),
    ("Автор", 22),
    ("Ответственный", 22),
    ("Приоритет", 14),
    ("Срок", 18),
    ("Статус", 22),
    ("Дата выполнения", 18),
    ("Дата закрытия", 18),
    ("Затраты", 14),
    ("Результат", 40),
    ("Просрочена", 12),
]


def _write_org_sheet(ws, requests) -> None:
    from ..db.models import OrgStatus

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="7030A0")
    for col, (name, width) in enumerate(ORG_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    for row_idx, req in enumerate(requests, start=2):
        values = [
            req.number,
            fmt_dt(req.created_at),
            req.object_name,
            req.brand.name if req.brand else "",
            req.type_title,
            req.description,
            req.quantity or "",
            req.author.full_name if req.author else "",
            req.assignee.full_name if req.assignee else "не назначен",
            req.priority_title,
            fmt_dt(req.due_at),
            OrgStatus.title(req.status),
            fmt_dt(req.done_at),
            fmt_dt(req.closed_at),
            float(req.cost or 0),
            req.result_comment or "",
            "да" if req.is_overdue else "нет",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            if col in (7, 16):
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def export_org_requests(requests, title: str = "Заявки УЦЦП") -> str:
    """Выгрузка организационных заявок — отдельный файл."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки УЦЦП"
    _write_org_sheet(ws, requests)

    os.makedirs(config.export_dir, exist_ok=True)
    stamp = now_local().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(config.export_dir, f"uccp_org_{stamp}.xlsx")
    wb.save(path)
    return path


def export_combined(tech_requests, org_requests, title: str = "Сводный отчёт") -> str:
    """Сводная выгрузка: два отдельных листа, статистика не смешивается."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ремонтные"
    _write_tech_sheet(ws, tech_requests)

    ws_org = wb.create_sheet("Организационные")
    _write_org_sheet(ws_org, org_requests)

    os.makedirs(config.export_dir, exist_ok=True)
    stamp = now_local().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(config.export_dir, f"uccp_combined_{stamp}.xlsx")
    wb.save(path)
    return path


async def run_export(func, *args, **kwargs) -> str:
    """Формирование файла в отдельном потоке — бот остаётся отзывчивым."""
    import asyncio

    return await asyncio.to_thread(func, *args, **kwargs)
